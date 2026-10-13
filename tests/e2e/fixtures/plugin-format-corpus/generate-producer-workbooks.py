from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

destination = Path(__file__).resolve().parent

workbook = Workbook()
sheet = workbook.active
sheet.title = "Produced budget"
sheet.append(["Producer-generated budget", "Amount", "Date", "Ratio"])
sheet.append(["Revenue – 世界", 12500.5, datetime(2026, 10, 8), 0.125])
sheet.append(["Expenses", 8750, True, "<script>literal content</script>"])
sheet.append(["Net", "=B2-B3"])
sheet["B2"].number_format = "0.00"
sheet["B3"].number_format = "0.00"
sheet["C2"].number_format = "yyyy-mm-dd"
sheet["D2"].number_format = "0.00%"
notes = workbook.create_sheet("Hidden notes")
notes["A1"] = "This hidden sheet is available in the picker"
notes.sheet_state = "hidden"
workbook.save(destination / "ProducedBudget.xlsx")

workbook = Workbook()
workbook.epoch = CALENDAR_MAC_1904
sheet = workbook.active
sheet.title = "1904 dates"
sheet.append([0, 1, 0.5])
sheet["A1"].number_format = "yyyy-mm-dd"
sheet["B1"].number_format = "yyyy-mm-dd"
sheet["C1"].number_format = "hh:mm:ss"
workbook.save(destination / "Produced1904.xlsx")

workbook = Workbook()
sheet = workbook.active
sheet.title = "Sparse limits"
sheet["XFD1000000"] = "Sparse boundary value"
workbook.save(destination / "ProducedSparse.xlsx")
